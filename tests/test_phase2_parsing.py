from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.ingestion.docx_parser import DocxParser
from app.ingestion.errors import DocumentParseError, OcrUnavailableError
from app.ingestion.ocr import TesseractOcr
from app.ingestion.pdf_parser import PdfParser, render_table_text


def _write_text_pdf(path: Path, page_texts: list[str]) -> Path:
    """Create a tiny local text PDF without an extra test dependency."""
    objects: list[bytes] = []
    page_ids = []
    for index in range(len(page_texts)):
        page_ids.append(4 + index * 2)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{value} 0 R" for value in page_ids).encode()
    objects.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(len(page_ids)).encode() + b" >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for index, text in enumerate(page_texts):
        content_id = page_ids[index] + 1
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("ascii")
        stream = b"BT /F1 12 Tf 72 720 Td (" + safe + b") Tj ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> /MediaBox [0 0 612 792] /Contents {content_id} 0 R >>".encode())
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    path.write_bytes(output)
    return path


def test_text_pdf_parsing_and_page_metadata(tmp_path: Path) -> None:
    path = _write_text_pdf(tmp_path / "sample.pdf", ["Guidance system technical description page one.", "Navigation system technical description page two."])
    parsed = PdfParser(None, min_text_characters=10, remove_margins=False).parse(path, "a" * 64)
    assert "Guidance system" in parsed.blocks[0].text
    assert {block.page_number for block in parsed.blocks} == {1, 2}
    assert not any(block.ocr_used for block in parsed.blocks)


def test_ocr_fallback_decision_is_page_threshold_based() -> None:
    parser = PdfParser(None, min_text_characters=10, min_alphanumeric_ratio=0.5)
    assert parser.needs_ocr("  12 ")
    assert parser.needs_ocr("---------- enough length ----------")
    assert not parser.needs_ocr("Technical INS content")


def test_empty_decorative_table_grid_is_dropped_but_real_rows_are_preserved() -> None:
    assert render_table_text([[None, "", None], ["", None, ""]]) == ""
    assert render_table_text([["Parameter", "Value"], ["Range", "150 km"]]) == (
        "Parameter | Value\nRange | 150 km"
    )


def test_broken_pdf_raises_safe_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a document and no confidential text")
    with pytest.raises(DocumentParseError, match="Invalid or unreadable PDF"):
        PdfParser(None).parse(path, "b" * 64)


def test_tesseract_missing_dependency_error_is_clear(monkeypatch: pytest.MonkeyPatch) -> None:
    import pytesseract
    monkeypatch.setattr(pytesseract, "get_languages", lambda config="": ["eng"])
    with pytest.raises(OcrUnavailableError, match="language data unavailable: tur"):
        TesseractOcr(("tur", "eng")).recognize(object())


def test_docx_preserves_heading_paragraph_and_table(tmp_path: Path) -> None:
    path = tmp_path / "technical.docx"
    document = Document()
    document.add_heading("Güdüm Sistemi", level=1)
    document.add_paragraph("INS ve GPS birlikte çalışır.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Parameter", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Range", "150 km"
    document.save(path)

    parsed = DocxParser().parse(path, "c" * 64)
    heading = next(block for block in parsed.blocks if block.kind == "heading")
    paragraph = next(block for block in parsed.blocks if block.kind == "paragraph")
    table_block = next(block for block in parsed.blocks if block.kind == "table")
    assert heading.heading_path == ("Güdüm Sistemi",)
    assert paragraph.section_title == "Güdüm Sistemi"
    assert paragraph.paragraph_index == 1
    assert paragraph.page_number is None
    assert "Parameter | Value" in table_block.text
    assert "Range | 150 km" in table_block.text
