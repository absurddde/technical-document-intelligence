"""Create a lightweight, local-only inventory of pilot PDF and DOCX files."""

from __future__ import annotations

import argparse
import csv
from dataclasses import asdict, dataclass
import hashlib
import json
from pathlib import Path
import re
import time
from typing import Any
import zipfile
from xml.etree import ElementTree

from app.ingestion.docx_parser import DocxParser
from app.ingestion.pdf_parser import PdfParser
from app.infrastructure.config import load_config


SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx"})
TURKISH_WORDS = frozenset({
    "ve", "bir", "bu", "ile", "için", "olarak", "olan", "sistem", "sistemler",
    "insansız", "araç", "araçlar", "askeri", "görev", "hedef", "teknoloji",
    "kullanılan", "tarafından", "üzerinde", "göre", "veya", "de", "da",
})
ENGLISH_WORDS = frozenset({
    "the", "and", "of", "to", "in", "for", "with", "is", "are", "system",
    "systems", "unmanned", "autonomous", "military", "technology", "vehicle",
    "vehicles", "from", "this", "that", "on", "or", "by", "as", "an",
})


@dataclass(slots=True)
class InventoryRecord:
    relative_path: str
    file_name: str
    extension: str
    size_bytes: int
    sha256: str
    pdf_page_count: int | None = None
    pdf_opened: bool | None = None
    encrypted: bool | None = None
    sampled_pages: list[int] | None = None
    sampled_native_characters: int | None = None
    native_chars_per_sampled_page: float | None = None
    native_text_page_ratio: float | None = None
    native_text_status: str | None = None
    probable_scanned: bool | None = None
    probable_language: str = "uncertain"
    sampled_table_count: int | None = None
    docx_paragraph_count: int | None = None
    docx_table_count: int | None = None
    parser_error: str | None = None
    notes: list[str] | None = None


def sha256_file(path: Path, block_size: int = 1024 * 1024) -> str:
    """Hash one file without loading it completely into memory."""

    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(block_size):
            digest.update(block)
    return digest.hexdigest()


def representative_indices(page_count: int, maximum: int = 5) -> list[int]:
    """Return evenly distributed zero-based page indices."""

    if page_count <= maximum:
        return list(range(page_count))
    return sorted({round(index * (page_count - 1) / (maximum - 1)) for index in range(maximum)})


def classify_language(text: str) -> str:
    """Estimate Turkish/English/mixed language from sampled native text."""

    words = re.findall(r"[^\W\d_]+", text.casefold(), re.UNICODE)
    if len(words) < 20:
        return "uncertain"
    turkish = sum(word in TURKISH_WORDS for word in words)
    english = sum(word in ENGLISH_WORDS for word in words)
    turkish_chars = len(re.findall(r"[çğıöşü]", text.casefold()))
    turkish_score = turkish + min(turkish_chars, 20) * 0.35
    if turkish_score >= 3 and english >= 3:
        weaker, stronger = sorted((turkish_score, float(english)))
        if weaker / stronger >= 0.25:
            return "mixed"
    if turkish_score >= 3 and turkish_score > english * 1.25:
        return "Turkish"
    if english >= 3 and english > turkish_score * 1.25:
        return "English"
    return "uncertain"


def repeated_margin_hint(page_texts: list[str]) -> bool:
    """Detect repeated first/last sampled lines without retaining their content."""

    counts: dict[str, int] = {}
    for text in page_texts:
        lines = [re.sub(r"\s+", " ", line).strip().casefold() for line in text.splitlines()]
        lines = [line for line in lines if 4 <= len(line) <= 160]
        for line in (lines[:1] + lines[-1:]):
            counts[line] = counts.get(line, 0) + 1
    return len(page_texts) >= 3 and any(count >= 3 for count in counts.values())


def possible_multicolumn(page: Any) -> bool:
    """Flag pages with many same-line words separated across both page halves."""

    words = page.extract_words() or []
    by_line: dict[int, list[dict[str, Any]]] = {}
    for word in words:
        by_line.setdefault(round(float(word["top"]) / 4), []).append(word)
    split_lines = 0
    for line_words in by_line.values():
        ordered = sorted(line_words, key=lambda item: float(item["x0"]))
        gaps = [float(right["x0"]) - float(left["x1"]) for left, right in zip(ordered, ordered[1:])]
        if gaps and max(gaps) > float(page.width) * 0.18:
            split_lines += 1
    return split_lines >= 4


def inspect_pdf(path: Path, record: InventoryRecord, parser: PdfParser) -> None:
    """Sample native PDF text and structural signals without OCR."""

    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            record.pdf_opened = True
            record.encrypted = bool(getattr(pdf.doc, "is_encrypted", False))
            record.pdf_page_count = len(pdf.pages)
            indices = representative_indices(record.pdf_page_count)
            record.sampled_pages = [index + 1 for index in indices]
            texts: list[str] = []
            sufficient = 0
            table_count = 0
            multicolumn_pages = 0
            for index in indices:
                page = pdf.pages[index]
                text = page.extract_text() or ""
                texts.append(text)
                sufficient += int(not parser.needs_ocr(text))
                try:
                    table_count += len(page.find_tables() or [])
                except Exception:
                    record.notes = [*(record.notes or []), f"table_detection_failed_page_{index + 1}"]
                try:
                    multicolumn_pages += int(possible_multicolumn(page))
                except Exception:
                    pass
            total_chars = sum(len(re.sub(r"\s+", "", text)) for text in texts)
            sample_count = len(indices)
            ratio = sufficient / sample_count if sample_count else 0.0
            record.sampled_native_characters = total_chars
            record.native_chars_per_sampled_page = round(total_chars / sample_count, 1) if sample_count else 0.0
            record.native_text_page_ratio = round(ratio, 3)
            record.sampled_table_count = table_count
            record.probable_language = classify_language("\n".join(texts))
            if ratio >= 0.8:
                record.native_text_status = "healthy"
                record.probable_scanned = False
            elif ratio <= 0.4:
                record.native_text_status = "low_or_absent"
                record.probable_scanned = True
            else:
                record.native_text_status = "mixed_or_uncertain"
                record.probable_scanned = False
                record.notes = [*(record.notes or []), "mixed_native_text_pages_review_before_ocr"]
            if repeated_margin_hint(texts):
                record.notes = [*(record.notes or []), "repeated_header_or_footer_observed"]
            if multicolumn_pages:
                record.notes = [*(record.notes or []), f"possible_multicolumn_on_{multicolumn_pages}_sampled_pages"]
            if table_count:
                record.notes = [*(record.notes or []), f"tables_detected_on_samples:{table_count}"]
            if record.pdf_page_count >= 150:
                record.notes = [*(record.notes or []), "very_large_page_count"]
    except Exception as error:
        record.pdf_opened = False
        record.parser_error = f"{type(error).__name__}: {error}"
        lowered = str(error).casefold()
        record.encrypted = "encrypt" in lowered or "password" in lowered


def inspect_docx(path: Path, record: InventoryRecord) -> None:
    """Inspect DOCX structure and exercise the existing project parser."""

    try:
        from docx import Document

        document = Document(path)
        record.docx_paragraph_count = len(document.paragraphs)
        record.docx_table_count = len(document.tables)
        parsed = DocxParser().parse(path, record.sha256)
        text = "\n".join(block.text for block in parsed.blocks)
        record.probable_language = classify_language(text)
        headings = sum(block.kind == "heading" for block in parsed.blocks)
        record.notes = [
            f"project_parser_blocks:{len(parsed.blocks)}",
            f"project_parser_headings:{headings}",
        ]
        if record.docx_table_count:
            record.notes.append(f"tables:{record.docx_table_count}")
    except (ImportError, OSError) as error:
        # The standard-library fallback keeps the probe useful if python-docx's
        # native lxml dependency is unavailable; it does not replace ingestion.
        try:
            with zipfile.ZipFile(path) as archive:
                root = ElementTree.fromstring(archive.read("word/document.xml"))
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            body = root.find(f"{namespace}body")
            children = list(body) if body is not None else []
            paragraphs = [child for child in children if child.tag == f"{namespace}p"]
            tables = [child for child in children if child.tag == f"{namespace}tbl"]
            text = "\n".join(
                "".join(node.text or "" for node in child.iter(f"{namespace}t"))
                for child in children
            )
            record.docx_paragraph_count = len(paragraphs)
            record.docx_table_count = len(tables)
            record.probable_language = classify_language(text)
            record.notes = [
                "docx_structure_read_with_stdlib_fallback",
                f"python_docx_unavailable:{type(error).__name__}",
            ]
            if tables:
                record.notes.append(f"tables:{len(tables)}")
        except Exception as fallback_error:
            record.parser_error = f"{type(fallback_error).__name__}: {fallback_error}"
    except Exception as error:
        record.parser_error = f"{type(error).__name__}: {error}"


def related_name_pairs(records: list[InventoryRecord]) -> list[dict[str, Any]]:
    """Identify obvious related/translated filename families without deduplicating them."""

    families = (
        ("ras_strategy", ("ras strategy", "robototic autonomous systems strategy")),
        ("autonomous_cooperative_levels", ("autonomous cooperative", "türkçe autonomous cooperative")),
    )
    pairs: list[dict[str, Any]] = []
    for family, phrases in families:
        matches = [record.file_name for record in records if any(
            phrase in re.sub(r"[_\-.]+", " ", record.file_name.casefold()) for phrase in phrases
        )]
        if len(matches) > 1:
            pairs.append({"family": family, "files": sorted(matches), "treat_as_duplicate": False})
    return pairs


def write_reports(report_dir: Path, report: dict[str, Any]) -> None:
    """Write local JSON and flattened CSV inventory reports."""

    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "pilot_inventory.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    records = report["documents"]
    if not records:
        return
    fieldnames = list(records[0])
    with (report_dir / "pilot_inventory.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        for item in records:
            flattened = dict(item)
            flattened["sampled_pages"] = ";".join(map(str, item["sampled_pages"] or []))
            flattened["notes"] = ";".join(item["notes"] or [])
            writer.writerow(flattened)


def probe(directory: Path, report_dir: Path) -> dict[str, Any]:
    """Probe supported documents under one supplied local directory."""

    root = directory.expanduser().resolve(strict=True)
    if not root.is_dir():
        raise ValueError(f"Not a directory: {root}")
    started = time.perf_counter()
    config = load_config(Path(__file__).resolve().parents[1] / "config" / "config.toml")
    pdf_parser = PdfParser(
        None,
        min_text_characters=config.parsing.pdf_min_text_characters,
        min_alphanumeric_ratio=config.parsing.pdf_min_alphanumeric_ratio,
    )
    records: list[InventoryRecord] = []
    files = sorted(
        (path for path in root.rglob("*") if path.is_file() and path.suffix.casefold() in SUPPORTED_EXTENSIONS),
        key=lambda path: str(path.relative_to(root)).casefold(),
    )
    for number, path in enumerate(files, 1):
        record = InventoryRecord(
            relative_path=str(path.relative_to(root)), file_name=path.name,
            extension=path.suffix.casefold(), size_bytes=path.stat().st_size,
            sha256=sha256_file(path), notes=[],
        )
        if record.extension == ".pdf":
            inspect_pdf(path, record, pdf_parser)
        else:
            inspect_docx(path, record)
        records.append(record)
        status = record.native_text_status or f"paragraphs={record.docx_paragraph_count}"
        print(f"[{number:02d}/{len(files):02d}] {record.file_name} | {status} | {record.probable_language}")

    hashes: dict[str, list[str]] = {}
    for record in records:
        hashes.setdefault(record.sha256, []).append(record.file_name)
    duplicate_groups = [sorted(names) for names in hashes.values() if len(names) > 1]
    summary = {
        "total_documents": len(records),
        "pdf_count": sum(record.extension == ".pdf" for record in records),
        "docx_count": sum(record.extension == ".docx" for record in records),
        "total_size_bytes": sum(record.size_bytes for record in records),
        "total_pdf_pages": sum(record.pdf_page_count or 0 for record in records),
        "healthy_native_text_pdfs": sum(record.native_text_status == "healthy" for record in records),
        "mixed_or_uncertain_native_text_pdfs": sum(record.native_text_status == "mixed_or_uncertain" for record in records),
        "probable_ocr_pdfs": sum(record.probable_scanned is True for record in records),
        "unreadable_files": sum(record.parser_error is not None for record in records),
        "languages": {
            language: sum(record.probable_language == language for record in records)
            for language in ("Turkish", "English", "mixed", "uncertain")
        },
        "duration_seconds": round(time.perf_counter() - started, 3),
    }
    report = {
        "probe_policy": "native text sampling only; no OCR, embeddings, FAISS, or LLM",
        "source_directory": str(root),
        "summary": summary,
        "exact_duplicate_groups": duplicate_groups,
        "related_name_families": related_name_pairs(records),
        "documents": [asdict(record) for record in records],
    }
    write_reports(report_dir.resolve(), report)
    print("\nSUMMARY")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path, help="Local directory containing PDF/DOCX files")
    parser.add_argument(
        "--report-dir", type=Path,
        default=Path(__file__).resolve().parents[1] / "data" / "pilot_reports",
    )
    arguments = parser.parse_args()
    probe(arguments.directory, arguments.report_dir)


if __name__ == "__main__":
    main()
